
from datetime import date
from pathlib import Path
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from book_manager.services.alerta_precio_service import Alerta


class ExportadorExcelService:
    """Genera un informe Excel con la comparación de precios."""

    @staticmethod
    def exportar(alertas: List[Alerta], carpeta: Path) -> Path:
        carpeta.mkdir(parents=True, exist_ok=True)

        libro = Workbook()
        hoja = libro.active
        hoja.title = "Comparación de precios"

        encabezados = [
            "Libro",
            "Precio interno",
            "Precio web",
            "Diferencia",
            "Fecha de extracción",
        ]
        hoja.append(encabezados)

        for alerta in alertas:
            hoja.append([
                alerta.titulo,
                alerta.precio_interno,
                alerta.precio_web,
                alerta.diferencia,
                alerta.fecha_extraccion,
            ])

        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill(
                start_color="4472C4",
                end_color="4472C4",
                fill_type="solid",
            )

        for columna in hoja.columns:
            hoja.column_dimensions[columna[0].column_letter].width = 22

        ruta = carpeta / f"comparacion_precios_{date.today().isoformat()}.xlsx"
        libro.save(ruta)

        return ruta
